"""Create conservative ATS documents from immutable approved evidence."""

import hashlib
import re
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .models import Job, Profile
from .policy import normalise, requirements

LABELS = {
    "skill": "Technical Skills",
    "project": "Projects",
    "experience": "Work Experience",
    "education": "Education",
    "language": "Languages",
    "award": "Awards",
}


def chronology(dates: str) -> tuple[int, int]:
    years = re.findall(r"\b(?:19|20)\d{2}\b", dates)
    months = [
        month
        for month in (
            "Jan",
            "Feb",
            "Mar",
            "Apr",
            "May",
            "Jun",
            "Jul",
            "Aug",
            "Sep",
            "Oct",
            "Nov",
            "Dec",
        )
        if month in dates[-12:]
    ]
    month_numbers = {
        month: index + 1
        for index, month in enumerate(
            ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
        )
    }
    return (int(years[-1]) if years else 0, month_numbers[months[-1]] if months else 0)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def font_name() -> str:
    for path in (Path("C:/Windows/Fonts/aptos.ttf"), Path("C:/Windows/Fonts/arial.ttf")):
        if path.is_file():
            if "CandidateFont" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("CandidateFont", str(path)))
                bold = path.parent / (
                    "arialbd.ttf" if path.name == "arial.ttf" else "aptos-bold.ttf"
                )
                if bold.is_file():
                    pdfmetrics.registerFont(TTFont("CandidateFont-Bold", str(bold)))
            return "CandidateFont"
    return "Helvetica"


def selected_lines(profile: Profile, job: Job, evidence_ids: list[str]) -> list[tuple[str, str]]:
    if not profile.confirmed or not profile.email or not profile.location:
        raise ValueError("Confirm the profile, e-mail and location before creating final documents")
    approved = {item.id: item for item in profile.evidence if item.verified}
    if not evidence_ids or any(eid not in approved for eid in evidence_ids):
        raise ValueError("Select at least one existing, verified evidence item")
    # Include factual history and education; choose skills/projects according to the job.
    selected = [
        item
        for item in profile.evidence
        if item.verified
        and (
            item.id in evidence_ids
            or item.category in {"experience", "education", "language", "award"}
        )
    ]
    lines = [
        ("name", profile.name),
        (
            "body",
            " | ".join(
                filter(
                    None,
                    (
                        profile.location,
                        profile.email,
                        profile.phone,
                    ),
                )
            ),
        ),
    ]
    lines.extend(("body", link) for link in profile.links)
    if profile.summary:
        lines.extend([("heading", "Professional Summary"), ("body", profile.summary)])
    needs = requirements(job)
    for category, heading in LABELS.items():
        items = [item for item in selected if item.category == category]
        items.sort(key=lambda item: chronology(item.dates), reverse=True)
        if items:
            lines.append(("heading", heading))
            for item in items:
                lines.append(
                    ("subheading", item.title + (f" | {item.dates}" if item.dates else ""))
                )
                if category == "skill" and item.tags:
                    relevant = [tag for tag in item.tags if normalise(tag) in needs]
                    if relevant:
                        lines.append(("body", ", ".join(relevant) + "."))
                    else:
                        lines.append(("body", item.text))
                else:
                    lines.append(("body", item.text))
    return lines


def write_pair(lines: list[tuple[str, str]], pdf: Path, docx: Path) -> int:
    font = font_name()
    bold_font = (
        "CandidateFont-Bold"
        if "CandidateFont-Bold" in pdfmetrics.getRegisteredFontNames()
        else "Helvetica-Bold"
    )
    styles = {
        kind: ParagraphStyle(
            kind,
            fontName=font if kind == "body" else bold_font,
            fontSize=size,
            leading=size * 1.25,
            alignment=TA_LEFT,
            textColor="black",
            spaceAfter=6,
            spaceBefore=8 if kind == "heading" else 0,
            keepWithNext=kind in {"heading", "subheading", "name"},
        )
        for kind, size in (("name", 19), ("heading", 13), ("subheading", 11), ("body", 11))
    }
    story: list[Any] = []
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    base = document.styles["Normal"]
    base.font.name, base.font.size = "Arial", Pt(11)
    base.font.color.rgb = RGBColor(0, 0, 0)
    for kind, text in lines:
        # Escape ReportLab markup: job descriptions cannot inject document instructions/links.
        clean = text.replace("\u2013", "-").replace("\u2014", "-")
        story.append(Paragraph(escape(clean).replace("\n", "<br/>"), styles[kind]))
        paragraph = document.add_paragraph()
        run = paragraph.add_run(clean)
        run.font.name, run.font.size = "Arial", Pt(styles[kind].fontSize)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.bold = kind in {"name", "heading", "subheading"}
        paragraph.paragraph_format.keep_with_next = kind != "body"
    story.append(Spacer(1, 1))
    SimpleDocTemplate(
        str(pdf),
        pagesize=A4,
        leftMargin=56.7,
        rightMargin=56.7,
        topMargin=56.7,
        bottomMargin=56.7,
        title=lines[0][1],
        author=lines[0][1],
    ).build(story)
    document.save(str(docx))
    reader = PdfReader(pdf)
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    if not text.strip() or lines[0][1] not in text or pdf.stat().st_size > 1_000_000:
        raise ValueError("Document text extraction or size validation failed")
    return len(reader.pages)


def generate(
    profile: Profile, job: Job, evidence_ids: list[str], folder: Path, revision: int
) -> dict[str, Any]:
    folder.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r"[^A-Za-z0-9_-]", "_", f"{profile.name}_{job.title}")[:110]
    files: dict[str, Any] = {}
    lines = selected_lines(profile, job, evidence_ids)
    cv_pdf, cv_docx = folder / f"{stem}_CV.pdf", folder / f"{stem}_CV.docx"
    pages = write_pair(lines, cv_pdf, cv_docx)
    if pages > 2:
        raise ValueError("CV exceeds two pages; reduce selected evidence before submission")
    pairs = {"cv_pdf": cv_pdf, "cv_docx": cv_docx}
    if job.cover_letter_required:
        approved = {item.id: item for item in profile.evidence if item.verified}
        letter = [
            ("name", profile.name),
            ("body", f"{profile.email} | {profile.location}"),
            ("heading", f"Application for {job.title} at {job.company}"),
            ("body", "Dear Hiring Team,"),
            (
                "body",
                f"I am applying for the {job.title} role. The following examples demonstrate my relevant experience and independent project work.",
            ),
        ]
        letter.extend(("body", approved[eid].text) for eid in evidence_ids[:3])
        letter.extend(
            [
                (
                    "body",
                    "I would welcome the opportunity to discuss how this experience relates to your team's work.",
                ),
                ("body", f"Yours faithfully,\n{profile.name}"),
            ]
        )
        letter_pdf, letter_docx = (
            folder / f"{stem}_Cover_Letter.pdf",
            folder / f"{stem}_Cover_Letter.docx",
        )
        if write_pair(letter, letter_pdf, letter_docx) > 2:
            raise ValueError("Cover letter is too long; select shorter evidence")
        pairs.update(cover_pdf=letter_pdf, cover_docx=letter_docx)
    for key, path in pairs.items():
        files[key] = {"name": path.name, "sha256": digest(path)}
    return {"revision": revision, "evidence_ids": evidence_ids, "pages": pages, "files": files}


def validate_manifest(manifest: dict[str, Any], folder: Path, revision: int) -> None:
    if manifest.get("revision") != revision or "cv_pdf" not in manifest.get("files", {}):
        raise ValueError("Generate documents for the current profile revision first")
    for item in manifest["files"].values():
        path = (folder / item["name"]).resolve()
        if path.parent != folder.resolve() or not path.is_file() or digest(path) != item["sha256"]:
            raise ValueError("Document missing, modified or outside the application directory")
